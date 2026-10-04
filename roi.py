"""Step 5 - ROI analysis: custom Hybrid-RAG build vs. third-party legal AI SaaS.

Implements the ROI template: baseline 2 hours manual review -> 20 minutes with the
agent; savings = time saved x volume x reviewer rate. It also compares the
one-time "build" cost of this custom RAG against the recurring "buy" cost of a
commercial legal-AI subscription, and can export the model to an Excel workbook -
the Build vs Buy decision at the heart of this use case.
"""
from __future__ import annotations

from dataclasses import dataclass


@dataclass
class ROIInputs:
    manual_minutes_per_doc: float = 120.0     # 2 hours baseline
    agent_minutes_per_doc: float = 20.0       # 20 minutes with the agent
    reviewer_hourly_rate: float = 75.0        # USD / hour
    documents_per_month: int = 200
    # Build (this custom hybrid-RAG)
    build_cost_one_time: float = 8000.0       # engineering build effort
    infra_cost_monthly: float = 50.0          # local / compute hosting
    # Buy (third-party legal AI SaaS)
    saas_platform_monthly: float = 1500.0
    saas_cost_per_doc: float = 8.0


def compute_roi(inp: ROIInputs) -> dict:
    minutes_saved = inp.manual_minutes_per_doc - inp.agent_minutes_per_doc
    hours_saved_per_doc = minutes_saved / 60.0
    cost_saved_per_doc = hours_saved_per_doc * inp.reviewer_hourly_rate

    monthly_savings = cost_saved_per_doc * inp.documents_per_month
    annual_savings = monthly_savings * 12

    build_year1 = inp.build_cost_one_time + inp.infra_cost_monthly * 12
    buy_year1 = (inp.saas_cost_per_doc * inp.documents_per_month
                 + inp.saas_platform_monthly) * 12

    net_year1 = annual_savings - build_year1
    roi_pct = (net_year1 / build_year1 * 100.0) if build_year1 else 0.0
    payback_months = (inp.build_cost_one_time / monthly_savings
                      if monthly_savings else float("inf"))
    build_vs_buy_annual = buy_year1 - build_year1

    return {
        "minutes_saved_per_doc": minutes_saved,
        "hours_saved_per_doc": hours_saved_per_doc,
        "cost_saved_per_doc": cost_saved_per_doc,
        "monthly_labor_savings": monthly_savings,
        "annual_labor_savings": annual_savings,
        "build_year1_cost": build_year1,
        "buy_year1_cost": buy_year1,
        "net_year1_savings_vs_manual": net_year1,
        "roi_pct_year1": roi_pct,
        "payback_period_months": payback_months,
        "build_vs_buy_annual_saving": build_vs_buy_annual,
    }


_INPUT_LABELS = {
    "manual_minutes_per_doc": "Manual review (min/doc)",
    "agent_minutes_per_doc": "Agent review (min/doc)",
    "reviewer_hourly_rate": "Reviewer rate (USD/hr)",
    "documents_per_month": "Documents / month",
    "build_cost_one_time": "Build cost (one-time, USD)",
    "infra_cost_monthly": "Infra cost (USD/month)",
    "saas_platform_monthly": "SaaS platform (USD/month)",
    "saas_cost_per_doc": "SaaS cost (USD/doc)",
}

_RESULT_LABELS = {
    "minutes_saved_per_doc": "Minutes saved / doc",
    "cost_saved_per_doc": "Cost saved / doc (USD)",
    "monthly_labor_savings": "Monthly labor savings (USD)",
    "annual_labor_savings": "Annual labor savings (USD)",
    "build_year1_cost": "Build - Year 1 cost (USD)",
    "buy_year1_cost": "Buy (SaaS) - Year 1 cost (USD)",
    "net_year1_savings_vs_manual": "Net Year 1 savings vs manual (USD)",
    "roi_pct_year1": "ROI % (Year 1)",
    "payback_period_months": "Payback period (months)",
    "build_vs_buy_annual_saving": "Build vs Buy annual saving (USD)",
}


def export_excel(inp: ROIInputs, results: dict, target) -> object:
    """Write the ROI model to an .xlsx file path or file-like buffer."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "ROI Model"

    bold = Font(bold=True)
    ws["A1"] = "Legal Document Review Agent - ROI Model"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "Inputs"
    ws["A3"].font = bold
    row = 4
    for key, label in _INPUT_LABELS.items():
        ws[f"A{row}"] = label
        ws[f"B{row}"] = getattr(inp, key)
        row += 1

    row += 1
    ws[f"A{row}"] = "Results"
    ws[f"A{row}"].font = bold
    row += 1
    for key, label in _RESULT_LABELS.items():
        ws[f"A{row}"] = label
        val = results.get(key, "")
        ws[f"B{row}"] = round(val, 2) if isinstance(val, (int, float)) else val
        row += 1

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18

    wb.save(target)
    return target


if __name__ == "__main__":
    inputs = ROIInputs()
    res = compute_roi(inputs)
    print("Legal Document Review Agent - ROI")
    for k, v in res.items():
        label = _RESULT_LABELS.get(k, k)
        print(f"  {label:34s}: {v:,.2f}" if isinstance(v, (int, float)) else f"  {label}: {v}")
